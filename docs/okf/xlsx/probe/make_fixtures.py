#!/usr/bin/env python3
"""Build the xlsx probe fixtures (openpyxl). Output: fixtures/*.xlsx.

Traps built into the data (each is what a clerk's real sheet looks like):
  attendance.xlsx          Attendance order != Roster order; blank punches; a punch exactly at shift start;
                           Late + no check-out; the shift start lives in Settings!B1, not in the formulas.
  attendance_status.xlsx   same day with Name/Dept/Hours/Status already typed in (summary task input).
  expenses.xlsx            23 ledger rows over 3 days + Budget sheet; Utilities has no spend.
  expenses_summary.xlsx    adds a finished Summary sheet (values) whose LAST row is a Total (chart range trap).
  ledger_plain.xlsx        ledger with dates stored as General serials (they display as 46321) and no styling.
  sales.xlsx               72 rows, Region x Month x Rep; no ties among rep totals.
  messy.xlsx               text dates in two styles, padded/mixed-case names, "ops"/"HR" case, text numbers,
                           blank rows, four duplicate rows once normalised.
  audit.xlsx               System Qty = 0 rows (divide by zero), a variance of exactly +2.0%.
  orders.xlsx              an order for an unknown customer, an order due exactly on the as-of date.
  daily_log.xlsx           chart + data validation + conditional format + merged title + defined name + Total row.
"""
import datetime as dt

import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.datetime import to_excel
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

import lib
from lib import *  # noqa: F401,F403

lib.FIX.mkdir(exist_ok=True)
HEAD_FONT = Font(bold=True, color="FFFFFF", name="Arial", size=10)
HEAD_PAT = PatternFill("solid", fgColor=HEAD_FILL)
BODY_FONT = Font(name="Arial", size=10)


def header(ws, names, row=1):
    for i, n in enumerate(names, 1):
        c = ws.cell(row, i, n)
        c.font, c.fill = HEAD_FONT, HEAD_PAT


def body(ws, first_row=2):
    for r in ws.iter_rows(min_row=first_row):
        for c in r:
            if c.value is not None and c.font == Font():
                c.font = BODY_FONT


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w


def attendance(with_status):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Attendance"
    names = {e: (n, d) for e, n, d in ROSTER}
    if with_status:
        header(ws, ["Date", "EmpID", "Name", "Dept", "Check-in", "Check-out", "Hours Worked", "Status"])
    else:
        header(ws, ["Date", "EmpID", "Check-in", "Check-out"])
    for emp in ATT_ORDER:
        i, o = PUNCH[emp]
        if with_status:
            ws.append([ATT_DATE, emp, names[emp][0], names[emp][1], hm(i), hm(o), att_hours(emp), att_status(emp)])
        else:
            ws.append([ATT_DATE, emp, hm(i), hm(o)])
    tcol = (5, 6) if with_status else (3, 4)
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 1).number_format = "yyyy-mm-dd"
        for c in tcol:
            ws.cell(r, c).number_format = "hh:mm"
    body(ws)
    widths(ws, [12, 8, 16, 10, 10, 10, 14, 14] if with_status else [12, 8, 10, 10])
    if not with_status:
        ro = wb.create_sheet("Roster")
        header(ro, ["EmpID", "Name", "Dept"])
        for r in ROSTER:
            ro.append(list(r))
        body(ro)
        widths(ro, [8, 18, 10])
        se = wb.create_sheet("Settings")
        se["A1"], se["B1"] = "Shift start", hm("09:00")
        se["B1"].number_format = "hh:mm"
        se["A1"].font = Font(bold=True)
    return wb


def ledger_sheet(wb, plain=False, title="Ledger"):
    ws = wb.active if wb.active.title == "Sheet" else wb.create_sheet()
    ws.title = title
    if plain:
        ws.append(LEDGER_HEAD)
        for d, cat, desc, amt, by, rc in LEDGER:
            ws.append([to_excel(d), cat, desc, amt, by, rc])      # General serial: shows as 46321
        return ws
    header(ws, LEDGER_HEAD)
    for row in LEDGER:
        ws.append(list(row))
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 1).number_format = "yyyy-mm-dd"
        ws.cell(r, 4).number_format = "#,##0.00"
    body(ws)
    widths(ws, [12, 16, 22, 10, 14, 8])
    return ws


def budget_sheet(wb):
    bs = wb.create_sheet("Budget")
    header(bs, ["Category", "Budget"])
    for r in BUDGET:
        bs.append(list(r))
    body(bs)
    widths(bs, [16, 10])


def expenses(summary=False, plain=False):
    wb = openpyxl.Workbook()
    ledger_sheet(wb, plain=plain)
    if plain:
        return wb
    budget_sheet(wb)
    if summary:
        ss = wb.create_sheet("Summary")
        header(ss, SUMMARY_HEAD)
        for cat, s, b, rem, pct, st in exp_summary():
            ss.append([cat, s, b, rem, pct, st])
        n = ss.max_row
        ss.append(["Total", round(sum(spent(c) for c, _ in BUDGET), 2), sum(b for _, b in BUDGET), None, None, None])
        ss.cell(n + 1, 4, round(ss.cell(n + 1, 3).value - ss.cell(n + 1, 2).value, 2))
        for r in range(2, n + 2):
            ss.cell(r, 5).number_format = "0.0%"
            if ss.cell(r, 5).value is None:
                ss.cell(r, 5, ss.cell(r, 2).value / ss.cell(r, 3).value)
            ss.cell(r, 5).number_format = "0.0%"
            for c in (2, 3, 4):
                ss.cell(r, c).number_format = "#,##0.00"
        for c in ss[n + 1]:
            c.font = Font(bold=True, name="Arial", size=10)
        body(ss)
        widths(ss, [16, 10, 10, 11, 9, 15])
    return wb


def sales():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    header(ws, ["Date", "Month", "Region", "Rep", "Product", "Amount"])
    for r in sales_rows():
        ws.append(list(r))
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 1).number_format = "yyyy-mm-dd"
        ws.cell(r, 6).number_format = "#,##0.00"
    body(ws)
    widths(ws, [12, 9, 9, 9, 10, 11])
    return wb


def messy():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Raw"
    header(ws, MESSY_HEAD)
    for r in MESSY:
        ws.append(list(r))
    for r in range(2, ws.max_row + 1):          # every text cell stays text; numbers stay numbers
        for c in (1, 4):
            if isinstance(ws.cell(r, c).value, str):
                ws.cell(r, c).number_format = "@"
    widths(ws, [14, 16, 10, 12])
    return wb


def audit():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Stock Count"
    header(ws, AUDIT_HEAD)
    for r in AUDIT:
        ws.append(list(r))
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 5).number_format = "#,##0.00"
    body(ws)
    widths(ws, [10, 20, 11, 12, 10])
    return wb


def orders():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Orders"
    header(ws, ["OrderID", "Order Date", "CustID", "Amount"])
    for r in ORDERS:
        ws.append(list(r))
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 2).number_format = "yyyy-mm-dd"
        ws.cell(r, 4).number_format = "#,##0.00"
    body(ws)
    widths(ws, [10, 12, 8, 10])
    cu = wb.create_sheet("Customers")
    header(cu, ["CustID", "Name", "Terms (days)"])
    for r in CUSTOMERS:
        cu.append(list(r))
    body(cu)
    widths(cu, [8, 20, 12])
    se = wb.create_sheet("Settings")
    se["A1"], se["B1"] = "As-of date", ASOF
    se["B1"].number_format = "yyyy-mm-dd"
    se["A1"].font = Font(bold=True)
    widths(se, [12, 12])
    return wb


def daily_log():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Log"
    ws["A1"] = LOG_TITLE
    ws.merge_cells("A1:D1")
    ws["A1"].font = Font(bold=True, size=14, name="Arial")
    ws["A1"].alignment = Alignment(horizontal="center")
    header(ws, LOG_HEAD, row=3)
    for i, row in enumerate(LOG, LOG_FIRST):
        for j, v in enumerate(row, 1):
            ws.cell(i, j, v)
        ws.cell(i, 1).number_format = "yyyy-mm-dd"
        ws.cell(i, 4).number_format = "#,##0.00"
    tr = LOG_LAST + 1
    ws.cell(tr, 1, "Total").font = Font(bold=True, name="Arial", size=10)
    ws.cell(tr, 4, f"=SUM(D{LOG_FIRST}:D{LOG_LAST})").font = Font(bold=True, name="Arial", size=10)
    ws.cell(tr, 4).number_format = "#,##0.00"
    body(ws, first_row=LOG_FIRST)
    dv = DataValidation(type="list", formula1=f'"{LOG_CATEGORIES}"', allow_blank=True)
    dv.add("C4:C20")
    ws.add_data_validation(dv)
    ws.conditional_formatting.add("D4:D20", CellIsRule(operator="greaterThan", formula=["100"],
                                                       fill=PatternFill("solid", bgColor="FFC7CE", fgColor="FFC7CE"),
                                                       font=Font(color="9C0006")))
    wb.defined_names["LogTotal"] = DefinedName("LogTotal", attr_text=f"Log!$D${tr}")
    ch = BarChart()
    ch.type, ch.title = "col", "Spend by item"
    ch.add_data(Reference(ws, min_col=4, min_row=LOG_FIRST, max_row=LOG_LAST), titles_from_data=False)
    ch.set_categories(Reference(ws, min_col=2, min_row=LOG_FIRST, max_row=LOG_LAST))
    ch.x_axis.delete = ch.y_axis.delete = False
    ch.legend = None
    ws.add_chart(ch, "F3")
    widths(ws, [12, 18, 12, 11])
    return wb


BUILD = {
    "attendance.xlsx": lambda: attendance(False),
    "attendance_status.xlsx": lambda: attendance(True),
    "expenses.xlsx": lambda: expenses(),
    "expenses_summary.xlsx": lambda: expenses(summary=True),
    "ledger_plain.xlsx": lambda: expenses(plain=True),
    "sales.xlsx": sales,
    "messy.xlsx": messy,
    "audit.xlsx": audit,
    "orders.xlsx": orders,
    "daily_log.xlsx": daily_log,
}

if __name__ == "__main__":
    for name, fn in BUILD.items():
        fn().save(lib.FIX / name)
        print("wrote", lib.FIX / name)
