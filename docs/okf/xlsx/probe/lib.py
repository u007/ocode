"""Shared data, expected answers and helpers for the xlsx probe.

make_fixtures.py, check.py and ref.py all import this module so the fixture, the
expected result and the reference solution can never drift apart.
"""
import contextlib
import datetime as dt
import os
import pathlib
import posixpath
import re
import shutil
import subprocess
import sys
import tempfile
import time
import zipfile

import openpyxl
from lxml import etree

HERE = pathlib.Path(__file__).resolve().parent
FIX = HERE / "fixtures"

NS = {
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
}

HEAD_FILL = "1F3B73"


def hm(s):
    """'08:52' -> datetime.time."""
    if s is None:
        return None
    h, m = s.split(":")
    return dt.time(int(h), int(m))


# ------------------------------------------------------------------ attendance

ROSTER = [
    ("E001", "Aisha Rahman", "Sales"), ("E002", "Ben Tan", "Sales"),
    ("E003", "Chloe Lim", "Ops"), ("E004", "Daniel Wong", "Ops"),
    ("E005", "Elena Cruz", "Finance"), ("E006", "Farid Hamid", "Finance"),
    ("E007", "Grace Lee", "HR"), ("E008", "Hakim Yusof", "HR"),
    ("E009", "Ivy Chen", "Sales"), ("E010", "Jon Ibrahim", "Ops"),
    ("E011", "Kavita Nair", "Finance"), ("E012", "Liam Ong", "HR"),
]
# EmpID -> (check-in, check-out). E004 punches exactly at the shift start (not late);
# E006/E010 never punched (a blank compares as 00:00, so a naive "< 09:00" test says on time);
# E007 forgot to check out; E012 is late AND has no check-out (Late wins).
PUNCH = {
    "E001": ("08:52", "17:31"), "E002": ("09:07", "17:45"), "E003": ("08:58", "17:02"),
    "E004": ("09:00", "17:30"), "E005": ("09:31", "18:10"), "E006": (None, None),
    "E007": ("08:45", None), "E008": ("08:30", "16:45"), "E009": ("09:15", "17:15"),
    "E010": (None, None), "E011": ("08:59", "17:00"), "E012": ("09:01", None),
}
ATT_DATE = dt.date(2026, 10, 28)
DEPTS = ["Sales", "Ops", "Finance", "HR"]
ATT_ORDER = [e for e, _, _ in reversed(ROSTER)]        # Attendance is NOT in Roster order


def att_status(emp, shift="09:00"):
    i, o = PUNCH[emp]
    if i is None:
        return "Absent"
    if hm(i) > hm(shift):
        return "Late"
    if o is None:
        return "No check-out"
    return "OK"


def att_hours(emp):
    i, o = PUNCH[emp]
    if i is None or o is None:
        return None
    a, b = hm(i), hm(o)
    return round(((b.hour * 60 + b.minute) - (a.hour * 60 + a.minute)) / 60, 2)


def att_summary(status_override=None):
    """dept -> (headcount, present, late, absent)."""
    out = {}
    for emp, _, dept in ROSTER:
        s = (status_override or {}).get(emp, att_status(emp))
        h, p, l, a = out.get(dept, (0, 0, 0, 0))
        out[dept] = (h + 1, p + (s != "Absent"), l + (s == "Late"), a + (s == "Absent"))
    return out


# -------------------------------------------------------------------- expenses

LEDGER_HEAD = ["Date", "Category", "Description", "Amount", "Paid By", "Receipt"]
LEDGER = [
    (dt.date(2026, 10, 26), "Travel", "Taxi to client site", 42.50, "Ben Tan", "Y"),
    (dt.date(2026, 10, 26), "Meals", "Team lunch", 128.40, "Aisha Rahman", "Y"),
    (dt.date(2026, 10, 26), "Office Supplies", "Printer paper", 64.90, "Chloe Lim", "Y"),
    (dt.date(2026, 10, 26), "Software", "Zoom licence", 79.00, "Elena Cruz", "Y"),
    (dt.date(2026, 10, 26), "Misc", "Courier", 23.00, "Ben Tan", "N"),
    (dt.date(2026, 10, 26), "Travel", "Parking", 18.00, "Daniel Wong", "Y"),
    (dt.date(2026, 10, 27), "Meals", "Client coffee", 36.80, "Aisha Rahman", "Y"),
    (dt.date(2026, 10, 27), "Travel", "Flight KL-SG", 412.00, "Ben Tan", "Y"),
    (dt.date(2026, 10, 27), "Office Supplies", "Toner", 189.00, "Chloe Lim", "Y"),
    (dt.date(2026, 10, 27), "Software", "Adobe CC", 220.00, "Elena Cruz", "Y"),
    (dt.date(2026, 10, 27), "Meals", "Working dinner", 96.20, "Daniel Wong", "Y"),
    (dt.date(2026, 10, 27), "Misc", "Gift basket", 55.00, "Grace Lee", "N"),
    (dt.date(2026, 10, 27), "Travel", "Taxi", 27.40, "Ben Tan", "Y"),
    (dt.date(2026, 10, 27), "Office Supplies", "Stapler", 15.60, "Chloe Lim", "Y"),
    (dt.date(2026, 10, 28), "Travel", "Hotel 1 night", 310.00, "Ben Tan", "Y"),
    (dt.date(2026, 10, 28), "Meals", "Breakfast meeting", 48.30, "Aisha Rahman", "Y"),
    (dt.date(2026, 10, 28), "Office Supplies", "Whiteboard markers", 22.40, "Chloe Lim", "Y"),
    (dt.date(2026, 10, 28), "Software", "Slack seats", 135.00, "Elena Cruz", "Y"),
    (dt.date(2026, 10, 28), "Misc", "Postage", 12.75, "Grace Lee", "Y"),
    (dt.date(2026, 10, 28), "Meals", "Team lunch", 142.10, "Aisha Rahman", "Y"),
    (dt.date(2026, 10, 28), "Travel", "Taxi", 31.60, "Daniel Wong", "Y"),
    (dt.date(2026, 10, 28), "Software", "Domain renewal", 18.50, "Elena Cruz", "N"),
    (dt.date(2026, 10, 28), "Misc", "Parcel insurance", 9.90, "Grace Lee", "Y"),
]
BUDGET = [("Travel", 800), ("Meals", 450), ("Office Supplies", 300),
          ("Utilities", 200), ("Software", 400), ("Misc", 80)]
SUMMARY_HEAD = ["Category", "Spent", "Budget", "Remaining", "% Used", "Status"]


def spent(cat, ledger=None):
    return round(sum(r[3] for r in (ledger or LEDGER) if r[1] == cat), 2)


def exp_summary(ledger=None):
    """Rows of the Summary sheet: (cat, spent, budget, remaining, pct_used, status)."""
    rows = []
    for cat, bud in BUDGET:
        s = spent(cat, ledger)
        rows.append((cat, s, bud, round(bud - s, 2), s / bud, "Over budget" if s > bud else "Within budget"))
    return rows


def day_stats(day, ledger=None):
    amts = [r[3] for r in (ledger or LEDGER) if r[0] == day]
    by_cat = {c: round(sum(r[3] for r in (ledger or LEDGER) if r[0] == day and r[1] == c), 2) for c, _ in BUDGET}
    return {"total": round(sum(amts), 2), "count": len(amts), "max": max(amts),
            "avg": round(sum(amts) / len(amts), 2), "by_cat": by_cat}


# ----------------------------------------------------------------------- sales

REGIONS = ["North", "South", "East", "West"]
MONTHS = ["2026-07", "2026-08", "2026-09"]
REPS = {"North": ["Amir", "Bella"], "South": ["Cheng", "Dewi"], "East": ["Emil", "Farah"], "West": ["Gita", "Hugo"]}
PRODUCTS = ["Alpha", "Beta", "Gamma"]


def sales_rows():
    """72 deterministic rows: (date, month, region, rep, product, amount)."""
    rows, x = [], 12345
    for mi, month in enumerate(MONTHS):
        for k in range(24):
            x = (x * 1103515245 + 12345) % (2 ** 31)
            region = REGIONS[k % 4]
            rep = REPS[region][(x >> 8) % 2]
            prod = PRODUCTS[(x >> 12) % 3]
            amount = 200 + ((x >> 4) % 4800) + ((x >> 20) % 100) / 100
            rows.append((dt.date(2026, 7 + mi, 1 + (x >> 16) % 27), month, region, rep, prod, round(amount, 2)))
    return rows


def sales_matrix(rows=None):
    rows = rows or sales_rows()
    return {r: {m: round(sum(x[5] for x in rows if x[2] == r and x[1] == m), 2) for m in MONTHS} for r in REGIONS}


def rep_totals(rows=None):
    rows = rows or sales_rows()
    tot = {}
    for x in rows:
        tot[x[3]] = round(tot.get(x[3], 0) + x[5], 2)
    return sorted(tot.items(), key=lambda kv: -kv[1])


# ----------------------------------------------------------------------- messy

MESSY_HEAD = ["Date", "Employee", "Dept", "Amount"]
MESSY = [
    ("2026-10-01", " alice tan ", "sales", "120.50"),
    ("2026-10-01", "BEN LIM", " Ops", 75),
    ("01 Oct 2026", "Alice Tan", "SALES ", "120.50"),         # duplicate of row 1
    ("2026-10-02", "chloe ong", "Finance", " 310.00 "),
    (None, None, None, None),
    ("2026-10-02", "Daniel  Wong", "ops", "45"),              # double space inside the name
    ("02 Oct 2026", "Chloe Ong", "finance", "310"),           # duplicate of row 4
    ("2026-10-03", "Elena Cruz", "HR", "88.25"),
    (None, None, None, None),
    ("2026-10-03", "ben lim", "OPS", "75"),                   # same person, different day than row 2
    ("2026-10-03", "Ben Lim", "Ops", 75),                     # duplicate of row 10
    ("03 Oct 2026", "elena cruz", "hr", "88.25"),             # duplicate of row 8
    ("2026-10-04", "Farid Hamid", "Finance", "1200"),
    ("2026-10-05", "grace lee ", "HR", "19.9"),
]
CLEAN = [
    (dt.date(2026, 10, 1), "Alice Tan", "Sales", 120.5),
    (dt.date(2026, 10, 1), "Ben Lim", "Ops", 75.0),
    (dt.date(2026, 10, 2), "Chloe Ong", "Finance", 310.0),
    (dt.date(2026, 10, 2), "Daniel Wong", "Ops", 45.0),
    (dt.date(2026, 10, 3), "Elena Cruz", "HR", 88.25),
    (dt.date(2026, 10, 3), "Ben Lim", "Ops", 75.0),
    (dt.date(2026, 10, 4), "Farid Hamid", "Finance", 1200.0),
    (dt.date(2026, 10, 5), "Grace Lee", "HR", 19.9),
]

# ----------------------------------------------------------------------- audit

AUDIT_HEAD = ["SKU", "Item", "System Qty", "Counted Qty", "Unit Cost"]
AUDIT = [
    ("SKU-001", "Copy paper A4", 120, 118, 4.50), ("SKU-002", "Toner black", 15, 12, 89.00),
    ("SKU-003", "Stapler", 40, 40, 6.20), ("SKU-004", "Pens blue", 500, 489, 0.45),
    ("SKU-005", "Folders", 200, 204, 1.10),            # +2.0% exactly: NOT flagged (strictly greater)
    ("SKU-006", "Whiteboard markers", 60, 55, 1.80), ("SKU-007", "Sticky notes", 0, 6, 2.20),
    ("SKU-008", "Binder clips", 300, 300, 0.30), ("SKU-009", "Envelopes", 1000, 1015, 0.12),
    ("SKU-010", "Lanyards", 80, 76, 0.90), ("SKU-011", "Cable ties", 0, 0, 0.05),
    ("SKU-012", "USB drives", 25, 24, 9.90), ("SKU-013", "Batteries AA", 144, 141, 0.75),
    ("SKU-014", "Label roll", 36, 36, 12.00), ("SKU-015", "Tape", 90, 92, 1.40),
    ("SKU-016", "Highlighters", 75, 74, 0.95),
]
AUDIT_SUMMARY_LABELS = ["Total system units", "Total counted units", "Net variance (units)",
                        "Items to investigate", "Net value impact"]


def audit_calc(rows=None):
    """Per row (variance, pct|None, value_impact, flag) plus the five summary figures."""
    out, flagged, nv = [], 0, 0.0
    for sku, item, sysq, cnt, cost in (rows or AUDIT):
        var = cnt - sysq
        pct = None if sysq == 0 else var / sysq
        flag = "Investigate" if ((sysq == 0 and cnt != 0) or (pct is not None and abs(pct) > 0.02)) else "OK"
        out.append((var, pct, round(var * cost, 2), flag))
        flagged += flag == "Investigate"
        nv += var * cost
    rr = rows or AUDIT
    summ = [sum(r[2] for r in rr), sum(r[3] for r in rr), sum(o[0] for o in out), flagged, round(nv, 2)]
    return out, summ


# ---------------------------------------------------------------------- orders

CUSTOMERS = [("C001", "Northwind Traders", 30), ("C002", "Contoso Ltd", 14), ("C003", "Fabrikam Inc", 45),
             ("C004", "Tailspin Toys", 7), ("C005", "Litware", 30)]
ORDERS = [
    ("O-1001", dt.date(2026, 9, 1), "C001", 1200), ("O-1002", dt.date(2026, 9, 15), "C002", 450),
    ("O-1003", dt.date(2026, 9, 20), "C003", 3100), ("O-1004", dt.date(2026, 10, 1), "C004", 85),
    ("O-1005", dt.date(2026, 10, 3), "C001", 640), ("O-1006", dt.date(2026, 10, 5), "C999", 220),
    ("O-1007", dt.date(2026, 10, 8), "C005", 910), ("O-1008", dt.date(2026, 10, 10), "C002", 75),
    ("O-1009", dt.date(2026, 10, 12), "C003", 1800), ("O-1010", dt.date(2026, 10, 15), "C004", 130),
    ("O-1011", dt.date(2026, 10, 20), "C001", 95), ("O-1012", dt.date(2026, 10, 23), "C005", 540),
    ("O-1013", dt.date(2026, 9, 30), "C002", 260), ("O-1014", dt.date(2026, 10, 29), "C004", 55),
    ("O-1015", dt.date(2026, 10, 16), "C002", 100),        # due exactly on the as-of date: not overdue
]
ASOF = dt.date(2026, 10, 30)


def orders_calc(asof=ASOF):
    """Per order (customer, due|None, days_overdue, status)."""
    cust = {c[0]: c for c in CUSTOMERS}
    out = []
    for oid, d, cid, amt in ORDERS:
        if cid not in cust:
            out.append(("UNKNOWN", None, 0, "Check customer"))
            continue
        due = d + dt.timedelta(days=cust[cid][2])
        late = max(0, (asof - due).days)
        out.append((cust[cid][1], due, late, "Overdue" if late > 0 else "Not due"))
    return out


# ------------------------------------------------------------------- daily log

LOG_TITLE = "Petty Cash Log - October 2026"
LOG_HEAD = ["Date", "Item", "Category", "Amount"]
LOG = [
    (dt.date(2026, 10, 23), "Printer ink", "Supplies", 64.00), (dt.date(2026, 10, 24), "Client lunch", "Meals", 120.00),
    (dt.date(2026, 10, 26), "Taxi", "Travel", 35.50), (dt.date(2026, 10, 27), "Stamps", "Misc", 12.00),
    (dt.date(2026, 10, 28), "Coffee beans", "Meals", 48.00), (dt.date(2026, 10, 28), "Parking", "Travel", 22.00),
]
LOG_CATEGORIES = "Travel,Meals,Supplies,Misc"
LOG_FIRST, LOG_LAST = 4, 9                 # data rows; Total sits on row 10
NEW_ENTRY = (dt.date(2026, 10, 29), "Courier fees", "Misc", 85.50)


# --------------------------------------------------------------------- helpers

def norm_date(v):
    """datetime / date / Excel serial -> date (or None)."""
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return (dt.datetime(1899, 12, 30) + dt.timedelta(days=float(v))).date()
    return None


class RecalcError(RuntimeError):
    """The recalculation backend could not run or produced nothing.

    This is an ENVIRONMENT failure (Excel missing, a permission prompt nobody answered, a
    timeout, ...), never a statement about the workbook the model saved. check.py reports it as
    ENV-FAIL so it is not scored as a model failure.
    """


# Excel's folder grant is per folder, so every recalculation is staged in this one directory:
# a single "Grant Access" click covers all later runs. The directory is intentionally left in
# place — an empty directory is harmless, and removing it would drop the grant.
STAGE_DIR = pathlib.Path.home() / "xlsx-probe-excel"
LOCK_PATH = pathlib.Path(tempfile.gettempdir()) / "xlsx-probe-excel.lock"
EXCEL_LOCK_TIMEOUT = 900.0          # a cold Excel start plus one unanswered prompt can take minutes
EXCEL_RUN_TIMEOUT = 180             # one recalculation, measured from process start


@contextlib.contextmanager
def _excel_lock(timeout=EXCEL_LOCK_TIMEOUT):
    """Serialise recalculations: every probe process shares the one Excel instance.

    Both platforms use the same contract — a non-blocking acquire in a retry loop with one
    deadline — so a Windows waiter does not fail after ~10 s while a macOS waiter blocks
    forever. Exceeding the deadline raises RecalcError (an environment failure)."""
    deadline = time.monotonic() + timeout
    with open(LOCK_PATH, "a+b") as fh:
        while True:
            try:
                if sys.platform == "win32":
                    import msvcrt
                    fh.seek(0)
                    msvcrt.locking(fh.fileno(), msvcrt.LK_NBLCK, 1)
                else:
                    import fcntl
                    fcntl.flock(fh, fcntl.LOCK_EX | fcntl.LOCK_NB)
            except OSError as exc:
                if time.monotonic() >= deadline:
                    raise RecalcError(
                        f"timed out after {timeout:.0f}s waiting for the Excel lock ({LOCK_PATH}): {exc}"
                    ) from exc
                time.sleep(0.2)
                continue
            break
        try:
            yield
        finally:
            if sys.platform == "win32":
                fh.seek(0)
                msvcrt.locking(fh.fileno(), msvcrt.LK_UNLCK, 1)
            else:
                fcntl.flock(fh, fcntl.LOCK_UN)


def _backend():
    """argv prefix (runner + script) for this platform, or raise RecalcError."""
    if sys.platform == "darwin":
        script = HERE / "excel_recalc.applescript"
        cmd = ["osascript", str(script)]
    elif sys.platform == "win32":
        script = HERE / "excel_recalc.ps1"
        cmd = ["powershell", "-NoProfile", "-File", str(script)]
    else:
        raise RecalcError(f"no Microsoft Excel backend for platform {sys.platform}")
    if not script.exists():
        raise RecalcError(f"missing recalculation script {script} (incomplete checkout?)")
    return cmd


def _remove_staged(*paths):
    for p in paths:
        try:
            p.unlink(missing_ok=True)
        except OSError as exc:
            # intentionally non-fatal: cleanup must not mask the real capture/recalc error,
            # but an operator still needs to know a staging file was left behind.
            print(f"warning: could not remove staging file {p}: {exc}", file=sys.stderr)


def recalc(path, outdir=None):
    """Recalculate through Microsoft Excel; return the path of the recalculated copy.

    macOS drives Excel with AppleScript, Windows through COM (PowerShell). Anywhere else there
    is no Excel, so it raises RecalcError (LibreOffice backend: TODO.md). The caller's file is
    never opened for writing: both the staged input and the processed output are copies in
    STAGE_DIR, so Excel cannot touch the artifact being graded."""
    cmd_prefix = _backend()  # fail before any filesystem change
    path = pathlib.Path(path).resolve()
    outdir = pathlib.Path(outdir or tempfile.mkdtemp(prefix="xlsx-recalc-")).resolve()
    outdir.mkdir(parents=True, exist_ok=True)
    out = outdir / path.name
    if out == path:
        raise ValueError(f"recalc output would overwrite its input: {path}")

    STAGE_DIR.mkdir(parents=True, exist_ok=True)
    # Unique per CALL, not per process: a workbook left open by an earlier failure can never be
    # reopened or clobbered by the next recalculation (check.py recalcs 2+ times per task).
    fd, staged = tempfile.mkstemp(prefix="in-", suffix=".xlsx", dir=STAGE_DIR)
    os.close(fd)
    src_in = pathlib.Path(staged)
    dst_in = src_in.with_name("out-" + src_in.name[len("in-"):])
    try:
        shutil.copyfile(path, src_in)
        if sys.platform == "win32":
            cmd = cmd_prefix + ["-Src", str(src_in), "-Dst", str(dst_in)]
        else:
            cmd = cmd_prefix + [str(src_in), str(dst_in)]
        try:
            with _excel_lock():
                res = subprocess.run(cmd, capture_output=True, text=True, errors="replace",
                                     timeout=EXCEL_RUN_TIMEOUT)
        except subprocess.TimeoutExpired as exc:
            raise RecalcError(f"Excel recalculation timed out after {exc.timeout}s for {path}") from exc
        if res.returncode != 0:
            detail = (res.stderr or res.stdout or "").strip()[:300] or f"exit {res.returncode}"
            raise RecalcError(f"Excel recalculation failed for {path}: {detail}")
        if not dst_in.exists():
            raise RecalcError(f"Excel produced no output for {path}")
        shutil.copyfile(dst_in, out)
    finally:
        _remove_staged(src_in, dst_in)
    if not out.exists():
        raise RecalcError(f"Excel produced no output for {path}")
    return out


def values(path):
    """Workbook of cached values of a recalculated copy (formulas become their results)."""
    return openpyxl.load_workbook(recalc(path), data_only=True)


def values_after_edit(path, edit):
    """Copy `path`, apply edit(wb) with openpyxl, recalc, return the data_only workbook."""
    tmp = pathlib.Path(tempfile.mkdtemp(prefix="xlsx-mut-"))
    wb = openpyxl.load_workbook(path)
    edit(wb)
    src = tmp / "mut.xlsx"
    wb.save(src)
    return values(src)


def header_map(ws, row=1):
    return {str(c.value).strip().lower(): c.column for c in ws[row] if c.value is not None}


# ---------------------------------------------------------------- chart parsing

def _rels(z, part):
    d, f = posixpath.split(part)
    rp = posixpath.join(d, "_rels", f + ".rels")
    if rp not in z.namelist():
        return {}
    out = {}
    for r in etree.fromstring(z.read(rp)).findall("rel:Relationship", NS):
        t = r.get("Target")
        t = t.lstrip("/") if t.startswith("/") else posixpath.normpath(posixpath.join(d, t))
        out[r.get("Id")] = (r.get("Type").rsplit("/", 1)[-1], t)
    return out


def parse_ref(f):
    """\"'Summary'!$A$2:$A$7\" -> ('Summary', 1, 2, 1, 7); single cell -> same corners."""
    m = re.match(r"^'?(.*?)'?!\$?([A-Z]+)\$?(\d+)(?::\$?([A-Z]+)\$?(\d+))?$", f or "")
    if not m:
        return None
    col = openpyxl.utils.column_index_from_string
    c2, r2 = (m.group(4) or m.group(2)), (m.group(5) or m.group(3))
    return (m.group(1), col(m.group(2)), int(m.group(3)), col(c2), int(r2))


def charts(path):
    """Every chart in the workbook: type, series refs, title, labels, axes, anchor, sheet."""
    out = []
    with zipfile.ZipFile(path) as z:
        wbx = etree.fromstring(z.read("xl/workbook.xml"))
        wrel = _rels(z, "xl/workbook.xml")
        for sh in wbx.find("m:sheets", NS):
            sheet_part = wrel[sh.get(f"{{{NS['r']}}}id")][1]
            for kind, dpart in _rels(z, sheet_part).values():
                if kind != "drawing":
                    continue
                drel = _rels(z, dpart)
                dx = etree.fromstring(z.read(dpart))
                for anc in dx:
                    frm = anc.find("xdr:from", NS)
                    ch = anc.find(".//c:chart", NS)
                    if ch is None:
                        continue
                    cpart = drel[ch.get(f"{{{NS['r']}}}id")][1]
                    info = chart_info(etree.fromstring(z.read(cpart)))
                    info.update(sheet=sh.get("name"), part=cpart,
                                col=int(frm.findtext("xdr:col", namespaces=NS)) if frm is not None else None,
                                row=int(frm.findtext("xdr:row", namespaces=NS)) if frm is not None else None)
                    out.append(info)
    return out


def chart_info(cx):
    plot = cx.find(".//c:plotArea", NS)
    kinds = [e for e in plot if e.tag.endswith("Chart")]
    first = kinds[0] if kinds else None
    ctype = etree.QName(first).localname if first is not None else None
    sers = []
    for s in (first.findall("c:ser", NS) if first is not None else []):
        name = s.find("c:tx", NS)
        sers.append({
            "name_ref": (name.findtext(".//c:f", namespaces=NS) if name is not None else None),
            "name_lit": (name.findtext(".//c:v", namespaces=NS) if name is not None else None),
            "cat": parse_ref(s.findtext("c:cat//c:f", namespaces=NS)),
            "val": parse_ref(s.findtext("c:val//c:f", namespaces=NS)),
        })
    title = cx.find("c:chart/c:title", NS)
    flags = lambda tag: [e.get("val") not in ("0", "false") for e in cx.iterfind(f".//c:dLbls/c:{tag}", NS)]
    axes = {}
    for ax in plot:
        n = etree.QName(ax).localname
        if n in ("catAx", "valAx", "dateAx"):
            d = ax.find("c:delete", NS)
            t = ax.find("c:title", NS)
            axes[n] = {"deleted": d is None or d.get("val") not in ("0", "false"),
                       "title": "".join(t.itertext()).strip() if t is not None else None}
    return {
        "type": ctype,
        "bar_dir": first.find("c:barDir", NS).get("val") if ctype == "barChart" else None,
        "grouping": (first.find("c:grouping", NS).get("val") if first is not None and first.find("c:grouping", NS) is not None else None),
        "series": sers,
        "title": "".join(title.itertext()).strip() if title is not None else None,
        "show_percent": any(flags("showPercent")),
        "show_value": any(flags("showVal")),
        "legend": cx.find("c:chart/c:legend", NS) is not None,
        "axes": axes,
    }
